"""Собирает шаблон Google Таблицы (xlsx) из CSV демо-каталога.
Загрузите файл в Google Диск → «Открыть с помощью Google Таблиц». Названия листов менять нельзя."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEETS = [('Товары', 'products'), ('Цены', 'prices'), ('Работы', 'works'), ('Комплектации', 'kits'), ('Поставщики', 'suppliers'), ('Расходники', 'consumables'), ('Категории', 'categories'), ('Объекты', 'object_types'), ('Позиции объектов', 'object_slots')]
UNITS = ['m2', 'm', 'pcs', 'set', 'kg', 'l', 'pack']
NUMERIC = {'pack_size', 'purchase_price', 'sale_price', 'price', 'delivery_price', 'delivery_days', 'rate', 'warranty_months', 'active', 'order', 'qty', 'waste', 'optional'}

wb = Workbook()
info = wb.active
info.title = 'Справка'
rows = [
    ['3D Home Academy — шаблон каталога'],
    ['ВНИМАНИЕ: значения в шаблоне — ДЕМО-числа для проверки расчётов, не рыночные цены. Замените их своими.'],
    [''],
    ['Листы, которые читает система (названия не менять):', ', '.join(s for s, _ in SHEETS)],
    ['Первая строка каждого листа — служебные названия колонок (не переводить и не переименовывать).'],
    ['product_id / work_id — постоянные идентификаторы. Не меняйте их у существующих позиций: на них ссылаются комплектации и заказы.'],
    ['Единицы измерения (unit):', ', '.join(UNITS) + '  (m2 — м², m — м, pcs — шт, set — комплект, kg — кг, l — литр, pack — упаковка)'],
    ['purchase_price — закупочная цена (видна только администратору); sale_price — продажная (идёт в смету).'],
    ['updated_at — дата обновления цены в формате ГГГГ-ММ-ДД.'],
    ['specs — характеристики «Название=значение;Название=значение»; attrs — служебные атрибуты (color, capacity_kw, series, protocol).'],
    ['Расходники: rate — норма; rate_per — на что умножается (m2, m, pcs или m2*mm); condition — условие, например screed=1.'],
    ['Подтверждённые заказы хранят цены на момент оформления — изменение таблицы их не меняет.'],
    [''],
    ['КОНСТРУКТОР: листы «Категории», «Объекты», «Позиции объектов».'],
    ['Категории — разделы левого меню (category_id, name, icon, order).'],
    ['Объекты — типы объектов квартиры: type_id, category_id, name, renderer (door|entrance|window|wall|none), placement (door|entrance|window|room_walls), primary_slot, description, lesson_steps (шаги через |), mistakes (через |).'],
    ['Позиции объектов — состав объекта: type_id, slot, label, qty (норма), per (pcs|opening_m2|opening_w|opening_h|opening_perimeter_m|wall_m2|floor_m2|perimeter_m|ceiling_m2), waste (запас, 0.1 = 10%), optional (1/0), default_product_id, work_id, work_per, order.'],
    ['Товар может переопределить норму позиции в attrs: rate=0.1;per=wall_m2;waste=0.15;work=W-WL-WP (например, обои — рулоны на м²).'],
    ['Новая модель двери или цвет краски = строка в «Товары» с тем же slot + строка в «Цены». Новый тип объекта = строка в «Объекты» + его позиции. Код не меняется.'],
]
for r in rows:
    info.append(r)
info['A1'].font = Font(bold=True, size=14)
info['A2'].font = Font(bold=True, color='9C5700')
info.column_dimensions['A'].width = 70
info.column_dimensions['B'].width = 90

head_fill = PatternFill('solid', fgColor='E8F0FF')
for title, name in SHEETS:
    ws = wb.create_sheet(title)
    with open(os.path.join(ROOT, 'data', 'catalog', f'{name}.csv'), encoding='utf-8') as f:
        data = list(csv.reader(f))
    header = data[0]
    for i, row in enumerate(data):
        out = []
        for col, v in zip(header, row):
            if i > 0 and col in NUMERIC and v.strip() != '':
                try:
                    v = float(v) if '.' in v else int(v)
                except ValueError:
                    pass
            out.append(v)
        ws.append(out)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
    ws.freeze_panes = 'A2'
    for idx, col in enumerate(header, start=1):
        width = max(len(str(r[idx - 1])) if idx - 1 < len(r) else 0 for r in data)
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = min(60, max(12, width + 2))
        ws.cell(1, idx).alignment = Alignment(wrap_text=False)
    def validate(colname, values):
        if colname in header:
            letter = ws.cell(1, header.index(colname) + 1).column_letter
            dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=False, showErrorMessage=True, error='Значение не из списка')
            ws.add_data_validation(dv)
            dv.add(f'{letter}2:{letter}2000')
    if name in ('products', 'works'):
        validate('unit', UNITS)
    if name == 'products':
        validate('category', ['material', 'equipment', 'consumable'])
    if name == 'kits':
        validate('tier', ['econom', 'comfort', 'premium'])

out = os.path.join(ROOT, 'docs', 'templates', '3d-home-academy-catalog-template.xlsx')
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print(out)
